#!/usr/bin/env python3
"""Build OAPGC genome long metadata from Nature Supplementary Data 2/5."""
import argparse, csv
from openpyxl import load_workbook


def rows(ws, header_row=3):
    it=ws.iter_rows(values_only=True)
    for _ in range(header_row-1): next(it)
    header=[str(x).strip() if x is not None else "" for x in next(it)]
    for values in it:
        yield {header[i]: values[i] if i<len(values) else None for i in range(len(header))}


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--xlsx",required=True)
    ap.add_argument("--annotation",required=True)
    ap.add_argument("--out",required=True)
    a=ap.parse_args()
    wb=load_workbook(a.xlsx,read_only=True,data_only=True)
    cluster={}
    awb=load_workbook(a.annotation,read_only=True,data_only=True)
    ws=awb[awb.sheetnames[0]]
    it=ws.iter_rows(values_only=True); header=[str(x).strip() if x is not None else "" for x in next(it)]
    ih=header.index("HQMAGs"); ic=header.index("SGBs"); ir=next(i for i,x in enumerate(header) if x.replace(" ","").lower()=="represent(strains)")
    for row in it:
        if not row or len(row)<=max(ih,ic,ir): continue
        gid,sgb,rep=row[ih],row[ic],row[ir]
        if gid and sgb and rep and str(gid)==str(rep): cluster[str(gid)]=str(sgb)
    sample={}
    for r in rows(wb["SD2"]):
        sid=r.get("Sample ID"); bio=r.get("NCBI BioSample ID"); typ=r.get("Sample type2")
        if sid: sample[str(sid)]=str(typ or "")
        if bio: sample[str(bio)]=str(typ or "")
    out=open(a.out,"w",newline="")
    w=csv.writer(out,delimiter="\t",lineterminator="\n")
    w.writerow(["genome_id","cluster","sample_id","biosample","site_raw","site","completeness","contamination","n50"])
    n=0
    for r in rows(wb["SD5"]):
        gid=str(r.get("Genomes name") or "")
        if not gid: continue
        sample_id=gid.split("_",1)[0]
        site_raw=sample.get(sample_id) or sample.get(gid) or ""
        w.writerow([gid,cluster.get(gid,""),sample_id,sample_id if sample_id.startswith("SAM") else "",site_raw,site_raw,
                    r.get("% Completeness"),r.get("% Contamination"),r.get("N50 length (bp)")])
        n+=1
    out.close()
    print(f"wrote {n} genomes")

if __name__=="__main__": main()
